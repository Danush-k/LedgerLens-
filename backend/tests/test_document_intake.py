"""Reading an FIR that arrives as a document.

Complaints reach an investigator as PDFs and Word files. Retyping a wallet
address out of one is the likeliest place for a transcription error to enter
a case, and a mistyped address traces a stranger's wallet with complete
confidence - so the extraction has to be exact, and its failures have to be
distinguishable from an empty result.
"""
import io

import pytest

from app.risk.complaint_parser import parse_complaint_text
from app.risk.document_text import UnreadableDocument, extract_text

SUSPECT = "1CRLGcaXajtWVF5EopZgQUqE12dKn8Rtuh"


def _pdf(text: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    y = 800
    for line in text.split("\n"):
        c.drawString(50, y, line)
        y -= 14
    c.save()
    return buf.getvalue()


def _docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    import docx

    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ── reading the file ──────────────────────────────────────────────────────

def test_reads_a_wallet_address_out_of_a_pdf():
    data = _pdf(f"FIR No. 0147/2026\nSuspect wallet: {SUSPECT}")

    text = extract_text("fir.pdf", data)

    assert SUSPECT in text


def test_reads_a_wallet_address_out_of_a_word_file():
    data = _docx([f"Suspect wallet address: {SUSPECT}"])

    assert SUSPECT in extract_text("fir.docx", data)


def test_reads_word_tables_not_just_paragraphs():
    """FIR forms are overwhelmingly tables. A paragraph-only read of one
    returns almost nothing, which would look like an empty complaint."""
    data = _docx(["First Information Report"],
                 table=[["Field", "Value"], ["Suspect Wallet", SUSPECT]])

    assert SUSPECT in extract_text("fir.docx", data)


def test_a_file_type_is_detected_from_content_not_only_the_name():
    """Uploads arrive with wrong or missing extensions."""
    data = _pdf(f"wallet {SUSPECT}")

    assert SUSPECT in extract_text("scan", data)


def test_plain_text_uploads_work():
    assert SUSPECT in extract_text("fir.txt", f"wallet {SUSPECT}".encode())


# ── failing honestly ──────────────────────────────────────────────────────

def test_a_scanned_pdf_says_it_has_no_text_layer():
    """"We could not read it" and "we read it and it holds nothing" are
    different findings, and only one of them means the complaint is empty."""
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    canvas.Canvas(buf, pagesize=A4).save()  # a page with no text at all

    with pytest.raises(UnreadableDocument) as exc:
        extract_text("scan.pdf", buf.getvalue())

    assert "scan" in str(exc.value).lower()


def test_an_unsupported_file_type_is_named_as_such():
    with pytest.raises(UnreadableDocument) as exc:
        extract_text("photo.png", b"\x89PNG\r\n\x1a\n binary")

    assert "pdf" in str(exc.value).lower()


def test_an_oversized_file_is_refused_before_parsing():
    with pytest.raises(UnreadableDocument) as exc:
        extract_text("huge.pdf", b"%PDF-" + b"0" * (11 * 1024 * 1024))

    assert "larger" in str(exc.value).lower()


def test_a_corrupt_pdf_is_reported_not_crashed():
    with pytest.raises(UnreadableDocument):
        extract_text("broken.pdf", b"%PDF-1.4 truncated garbage")


# ── the whole path ────────────────────────────────────────────────────────

def test_document_to_parsed_wallet_end_to_end():
    data = _pdf(
        "FIRST INFORMATION REPORT\n"
        "FIR No. 0147/2026\n"
        "NCRP Acknowledgement No. NCRP/2026/KA/0004471\n"
        f"Suspect Wallet Address (BTC): {SUSPECT}\n"
        "Amount defrauded: Rs. 47,50,000"
    )

    parsed = parse_complaint_text(extract_text("fir.pdf", data))

    assert [w["address"] for w in parsed["wallets"]] == [SUSPECT]
    assert parsed["wallets"][0]["chain"] == "bitcoin"
    assert any("0147/2026" in r for r in parsed["complaint_refs"])


def test_references_do_not_capture_surrounding_words():
    """The regression: the pattern matched the keyword alone, so a real FIR
    yielded "FIRST" and "FIR PARTICULARS" as complaint references. A field
    auto-filled with the word "FIRST" is worse than a blank one, because it
    looks filled in."""
    parsed = parse_complaint_text(
        "FIRST INFORMATION REPORT\nDISTRICT / FIR PARTICULARS\nFIR No. 0147/2026")

    assert all(any(ch.isdigit() for ch in r) for r in parsed["complaint_refs"])
    assert "FIRST" not in parsed["complaint_refs"]
