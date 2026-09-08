"""Pulling plain text out of an uploaded FIR.

Complaints reach an investigator as documents, not as pasted text: a scanned
FIR, a Word file from the complainant, a PDF export from NCRP. Making them
retype a wallet address out of one is both slow and the single most likely
place for a transcription error to enter the case - and a mistyped address
traces a stranger's wallet with full confidence.

Only the text layer is read. A scanned FIR that is pure images will yield
nothing, and this says so plainly rather than returning an empty string that
the caller would report as "no wallets found" - the difference between "we
could not read it" and "we read it and it holds nothing" matters here as
much as it does anywhere else in this system.
"""
import io
import re


class UnreadableDocument(Exception):
    """Raised when the file cannot be read as text at all."""


MAX_BYTES = 10 * 1024 * 1024


def extract_text(filename: str, data: bytes) -> str:
    """Best-effort plain text from a PDF, Word or plain-text upload."""
    if len(data) > MAX_BYTES:
        raise UnreadableDocument(
            f"File is larger than {MAX_BYTES // (1024 * 1024)}MB.")

    name = (filename or "").lower()
    if name.endswith(".pdf") or data[:5] == b"%PDF-":
        return _clean(_from_pdf(data))
    if name.endswith((".docx", ".doc")) or data[:2] == b"PK":
        return _clean(_from_docx(data))
    try:
        return _clean(data.decode("utf-8"))
    except UnicodeDecodeError:
        raise UnreadableDocument(
            "Unsupported file type. Upload a PDF, Word (.docx) or plain-text file.")


def _from_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:  # pragma: no cover - dependency is declared
        raise UnreadableDocument("PDF support is not installed on the server.")

    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:  # noqa: BLE001 - any parse failure is unreadable
        raise UnreadableDocument(f"Could not read the PDF ({type(exc).__name__}).")

    text = "\n".join(pages)
    if not text.strip():
        # Almost always a scan. Saying so points at the fix - paste the text,
        # or run OCR - instead of implying the document held nothing.
        raise UnreadableDocument(
            "This PDF has no text layer - it is most likely a scan. "
            "Paste the details as text, or upload a text-based PDF.")
    return text


def _from_docx(data: bytes) -> str:
    try:
        import docx
    except ImportError:  # pragma: no cover - dependency is declared
        raise UnreadableDocument("Word support is not installed on the server.")

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise UnreadableDocument(f"Could not read the Word file ({type(exc).__name__}).")

    parts = [p.text for p in document.paragraphs]
    # FIR forms are overwhelmingly tables - a paragraph-only read of one
    # returns almost nothing, which would look like an empty complaint.
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)

    text = "\n".join(parts)
    if not text.strip():
        raise UnreadableDocument("The Word file contains no readable text.")
    return text


def _clean(text: str) -> str:
    """Normalise whitespace without joining separate fields together.

    PDF extraction often breaks a wallet address across a line. Runs of
    spaces are collapsed, but line structure is kept, because an FIR's
    meaning lives in its field-per-line layout.
    """
    text = text.replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()
